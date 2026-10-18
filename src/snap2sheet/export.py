"""CSV and Excel export."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Receipt

COLUMNS = [
    "file",
    "merchant",
    "date",
    "total",
    "tax",
    "currency",
    "invoice_number",
    "items",
    "confidence",
    "warnings",
]


def write_csv(receipts: list[Receipt], path: Path | str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=COLUMNS)
        writer.writeheader()
        for receipt in receipts:
            writer.writerow(receipt.to_row())
    return path


def rows_to_xlsx(rows: list[dict[str, object]], path: Path | str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "expenses"
    sheet.append(COLUMNS)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F7A4D")
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
    for row in rows:
        sheet.append([row.get(column, "") for column in COLUMNS])
    widths = {"file": 28, "merchant": 30, "date": 12, "total": 10, "tax": 10,
              "currency": 9, "invoice_number": 16, "items": 7, "confidence": 11, "warnings": 45}
    for index, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(column, 14)
    sheet.freeze_panes = "A2"
    workbook.save(path)
    return path


def write_xlsx(receipts: list[Receipt], path: Path | str) -> Path:
    return rows_to_xlsx([receipt.to_row() for receipt in receipts], path)
