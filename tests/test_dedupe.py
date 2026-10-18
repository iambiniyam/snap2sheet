import csv

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from snap2sheet.cli import app
from snap2sheet.dedupe import (
    annotate_duplicates,
    count_duplicates,
    duplicate_key,
    normalize_date,
    normalize_merchant,
    normalize_total,
)
from snap2sheet.models import Receipt


def test_exact_duplicates_are_marked_and_kept():
    receipts = [
        Receipt(
            source="a.png", merchant="Blue Bottle Coffee", date="2026-03-12", total=14.58
        ),
        Receipt(
            source="b.png", merchant="  BLUE   BOTTLE COFFEE! ", date="2026/03/12", total=14.580
        ),
    ]

    assert len({duplicate_key(receipt) for receipt in receipts}) == 1

    annotated = annotate_duplicates(receipts)
    assert len(annotated) == 2
    assert annotated[0].duplicate is False
    assert annotated[0].duplicate_of is None
    assert annotated[1].duplicate is True
    assert annotated[1].duplicate_of == "a.png"
    assert count_duplicates(annotated) == 1


def test_same_merchant_and_total_on_different_dates_is_not_duplicate():
    receipts = [
        Receipt(source="a.png", merchant="Cafe", date="2026-03-12", total=10.0),
        Receipt(source="b.png", merchant="Cafe", date="2026-03-13", total=10.0),
    ]

    annotate_duplicates(receipts)

    assert len({duplicate_key(receipt) for receipt in receipts}) == 2
    assert all(receipt.duplicate is False for receipt in receipts)
    assert all(receipt.duplicate_of is None for receipt in receipts)


@pytest.mark.parametrize("missing", ["merchant", "date", "total"])
def test_missing_fields_are_never_reported_as_duplicates(missing):
    values = {"merchant": "Cafe", "date": "2026-03-12", "total": 10.0}
    values[missing] = None
    receipts = [
        Receipt(source="a.png", **values),
        Receipt(source="b.png", **values),
    ]

    assert all(duplicate_key(receipt) is None for receipt in receipts)
    annotate_duplicates(receipts)
    assert all(receipt.duplicate is False for receipt in receipts)
    assert all(receipt.duplicate_of is None for receipt in receipts)


def test_ambiguous_and_invalid_dates_are_conservative():
    assert normalize_date("03/04/2026") is None
    assert normalize_date("2026-02-30") is None
    assert normalize_date("not a date") is None
    assert duplicate_key(
        Receipt(source="a.png", merchant="Cafe", date="03/04/2026", total=10.0)
    ) is None


def test_normalization_helpers():
    assert normalize_merchant("  Café  No. 1 ") == "café no 1"
    assert normalize_date("2026-3-4") == "2026-03-04"
    assert normalize_date("12 Mar 2026") == "2026-03-12"
    assert normalize_total(14.5) == "14.50"
    assert normalize_total("14.58") == "14.58"
    assert normalize_total(float("nan")) is None
    assert normalize_total("1,234.56") is None


def test_annotation_clears_stale_flags():
    receipt = Receipt(
        source="a.png",
        merchant="Cafe",
        date="2026-03-12",
        total=5.0,
        duplicate=True,
        duplicate_of="old.png",
    )

    annotate_duplicates([receipt])

    assert receipt.duplicate is False
    assert receipt.duplicate_of is None


def test_receipt_dict_and_row_expose_duplicate_flags():
    original = Receipt(source="a.png", merchant="Cafe", date="2026-03-12", total=10.0)
    duplicate = Receipt(source="b.png", merchant="Cafe", date="2026-03-12", total=10.0)
    annotate_duplicates([original, duplicate])

    payload = duplicate.to_dict()
    assert payload["duplicate"] is True
    assert payload["duplicate_of"] == "a.png"

    row = duplicate.to_row()
    assert row["duplicate"] == "yes"
    assert row["duplicate_of"] == "a.png"


def test_batch_keeps_and_marks_duplicates(tmp_path, monkeypatch):
    folder = tmp_path / "receipts"
    folder.mkdir()
    (folder / "a.png").write_bytes(b"a")
    (folder / "b.png").write_bytes(b"b")
    receipts = {
        "a.png": Receipt(source="a.png", merchant="Cafe", date="2026-03-12", total=10.0),
        "b.png": Receipt(source="b.png", merchant="CAFE", date="2026-03-12", total=10.00),
    }

    def fake_scan(path, order, use_llm, ocr="auto"):
        return receipts[path.name]

    monkeypatch.setattr("snap2sheet.cli._scan", fake_scan)
    output = tmp_path / "expenses.csv"
    xlsx = tmp_path / "expenses.xlsx"

    result = CliRunner().invoke(
        app, ["batch", str(folder), "--output", str(output), "--xlsx", str(xlsx)]
    )

    assert result.exit_code == 0, result.output
    with output.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    assert [row["file"] for row in rows] == ["a.png", "b.png"]
    assert rows[0]["duplicate"] == ""
    assert rows[1]["duplicate"] == "yes"
    assert rows[1]["duplicate_of"] == "a.png"
    assert "possible duplicate" in result.output

    sheet = load_workbook(xlsx).active
    header = [cell.value for cell in sheet[1]]
    assert "duplicate" in header
    assert "duplicate_of" in header
    assert sheet.cell(row=3, column=header.index("duplicate") + 1).value == "yes"
