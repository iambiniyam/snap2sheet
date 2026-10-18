from openpyxl import load_workbook

from snap2sheet.export import rows_to_xlsx, safe_cell, write_csv
from snap2sheet.models import Receipt


def test_safe_cell_blocks_formulas():
    assert safe_cell("=HYPERLINK(\"http://x\")").startswith("'=")
    assert safe_cell("+1") == "'+1"
    assert safe_cell("-1") == "'-1"
    assert safe_cell(14.58) == 14.58


def test_xlsx_cells_are_not_formulas(tmp_path):
    path = rows_to_xlsx(
        [{"merchant": "=HYPERLINK(\"http://x\")", "total": "10.00"}],
        tmp_path / "out.xlsx",
    )
    sheet = load_workbook(path, data_only=False).active
    merchant = sheet.cell(row=2, column=2).value
    assert merchant.startswith("'=")
    assert sheet.cell(row=2, column=1).value in (None, "")


def test_csv_cells_are_not_formulas(tmp_path):
    receipt = Receipt(source="r.png", merchant="=cmd|' /C calc'!A0")
    path = write_csv([receipt], tmp_path / "out.csv")
    text = path.read_text()
    assert "'=cmd" in text
    assert "\n=cmd" not in text
